import io
import docx
from app.ingestion.extractors.base import BaseExtractor, ExtractedDocument, ExtractedPage


class DocxExtractor(BaseExtractor):
    def extract(self, content: bytes) -> ExtractedDocument:
        try:
            doc_file = io.BytesIO(content)
            doc = docx.Document(doc_file)
        except Exception as e:
            raise ValueError(f"Unable to parse DOCX document: {str(e)}") from e

        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs).strip()
        if not full_text:
            raise ValueError("DOCX document contains no readable text.")

        page = ExtractedPage(page_number=None, text=full_text)
        return ExtractedDocument(
            pages=[page],
            page_count=None,
            total_characters=len(full_text)
        )
